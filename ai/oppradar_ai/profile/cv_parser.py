"""Parse a CV (PDF / DOCX / TXT) into a structured Profile.

Two layers:
1. Deterministic text extraction per file type (offline, always available).
2. Optional LLM structuring — sections, skills, education, experience.

Works without an LLM key via a keyword-based fallback (lower quality but
functional); raises CVParseError only when the file itself can't be read.
"""

from __future__ import annotations

import json
import logging
import re
from pathlib import Path

from oppradar_ai import llm
from oppradar_ai.errors import CVParseError
from oppradar_ai.schemas import Education, Experience, Profile

logger = logging.getLogger(__name__)

SUPPORTED_SUFFIXES = {".pdf", ".docx", ".txt", ".md"}

# Section heading detection for the fallback parser.
SECTION_PATTERNS: dict[str, re.Pattern[str]] = {
    "summary": re.compile(r"^\s*(profile|summary|about|à propos)\b", re.I),
    "experience": re.compile(r"^\s*(work\s+)?(experience|professional\s+experience|parcours)\b", re.I),
    "education": re.compile(r"^\s*(education|formation|academic|qualifications?)\b", re.I),
    "skills": re.compile(r"^\s*(skills|compétences|technical\s+skills|abilities)\b", re.I),
    "languages": re.compile(r"^\s*(languages?|langues?)\b", re.I),
}

# Common CV skill vocabulary — the fallback extractor leans on this list.
SKILL_VOCAB = [
    "python", "javascript", "typescript", "java", "c++", "sql", "html", "css",
    "react", "next.js", "node.js", "django", "fastapi", "flask", "spring",
    "machine learning", "deep learning", "nlp", "computer vision", "pytorch",
    "tensorflow", "pandas", "numpy", "docker", "kubernetes", "aws", "gcp",
    "azure", "git", "linux", "figma", "excel", "power bi", "tableau",
    "communication", "leadership", "project management", "agile",
]

LANG_VOCAB = ["arabic", "franchais", "français", "french", "english", "anglais", "german", "allemand", "spanish", "espagnol"]


# --------------------------------------------------------------------------- #
# Text extraction
# --------------------------------------------------------------------------- #
def _extract_text(path: Path) -> str:
    suffix = path.suffix.lower()

    if suffix in {".txt", ".md"}:
        return path.read_text(encoding="utf-8", errors="ignore")

    if suffix == ".pdf":
        try:
            from pypdf import PdfReader
        except ImportError as exc:  # pragma: no cover
            raise CVParseError("pypdf is required for PDF CVs: pip install pypdf") from exc
        try:
            reader = PdfReader(str(path))
            text = "\n".join(page.extract_text() or "" for page in reader.pages)
        except Exception as exc:
            raise CVParseError(f"could not read PDF {path.name}: {exc}") from exc
        if not text.strip():
            raise CVParseError(f"{path.name} contains no extractable text (scanned image?)")
        return text

    if suffix == ".docx":
        try:
            from docx import Document
        except ImportError as exc:  # pragma: no cover
            raise CVParseError("python-docx is required for DOCX CVs: pip install python-docx") from exc
        try:
            document = Document(str(path))
        except Exception as exc:
            raise CVParseError(f"could not read DOCX {path.name}: {exc}") from exc
        paragraphs = [p.text for p in document.paragraphs]
        tables = [
            cell.text for table in document.tables for row in table.rows for cell in row.cells
        ]
        return "\n".join(paragraphs + tables)

    raise CVParseError(f"unsupported file type {suffix!r} (expected one of {sorted(SUPPORTED_SUFFIXES)})")


# --------------------------------------------------------------------------- #
# LLM structuring
# --------------------------------------------------------------------------- #
PROMPT = """You structure a CV into JSON. Return ONLY a JSON object:
{{
  "full_name": str,
  "headline": str,
  "summary": str,
  "skills": [str, ...],
  "languages": [str, ...],
  "interests": [str, ...],
  "education": [{{"institution": str, "degree": str, "field": str,
                  "start_date": str|null, "end_date": str|null}}, ...],
  "experience": [{{"organization": str, "title": str, "start_date": str|null,
                   "end_date": str|null, "description": str}}, ...]
}}
Normalize skills to lowercase canonical names (e.g. "JS" -> "javascript").
Never invent entries that are not in the CV.

CV text:
---
{cv_text}
---
"""


def _structure_with_llm(cv_text: str) -> Profile:
    data = llm.chat_json([{"role": "user", "content": PROMPT.format(cv_text=cv_text[:6000])}])
    data["education"] = [Education.model_validate(e) for e in data.get("education", [])]
    data["experience"] = [Experience.model_validate(e) for e in data.get("experience", [])]
    return Profile.model_validate(data)


# --------------------------------------------------------------------------- #
# Keyword fallback (no LLM)
# --------------------------------------------------------------------------- #
def _structure_fallback(cv_text: str) -> Profile:
    lines = [ln.strip() for ln in cv_text.splitlines()]
    sections: dict[str, list[str]] = {key: [] for key in SECTION_PATTERNS}
    current: str | None = None
    for line in lines:
        matched = next((key for key, pat in SECTION_PATTERNS.items() if pat.match(line)), None)
        if matched:
            current = matched
            continue
        if current and line:
            sections[current].append(line)

    lowered = cv_text.lower()
    skills = [s for s in SKILL_VOCAB if s in lowered]
    languages = sorted({l for l in LANG_VOCAB if l in lowered})

    name = next((ln for ln in lines if ln and len(ln.split()) <= 4 and not ln.isupper()), "")
    experience_lines = sections["experience"][:5]
    return Profile(
        full_name=name,
        skills=skills,
        languages=languages,
        summary=" ".join(sections["summary"][:3]),
        experience=[Experience(title=ln[:80]) for ln in experience_lines],
        education=[Education(degree=ln[:80]) for ln in sections["education"][:4]],
    )


# --------------------------------------------------------------------------- #
# Public API
# --------------------------------------------------------------------------- #
def parse_cv(path: str | Path, *, use_llm: bool = True) -> Profile:
    """Parse a CV file into a structured Profile (interface #2 in docs/ai-interface.md)."""
    path = Path(path)
    if not path.exists():
        raise CVParseError(f"file not found: {path}")
    if path.stat().st_size == 0:
        raise CVParseError(f"{path.name} is empty")

    cv_text = _extract_text(path)
    if not cv_text.strip():
        raise CVParseError(f"{path.name} contains no extractable text")

    if use_llm and llm.is_configured():
        try:
            return _structure_with_llm(cv_text)
        except Exception as exc:
            logger.warning("LLM CV structuring failed, using fallback: %s", exc)

    return _structure_fallback(cv_text)
