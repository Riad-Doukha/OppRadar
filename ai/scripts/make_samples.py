# Generates the SYNTHETIC sample documents (CVs, certificate, portfolio) and
# their ground-truth JSON under ai/data/samples/. No real personal data.
# Run from ai/:  python scripts/make_samples.py

from __future__ import annotations

import json
from pathlib import Path

import pymupdf
from docx import Document

SAMPLES = Path(__file__).resolve().parents[1] / "data" / "samples"
GROUND_TRUTH = SAMPLES / "ground_truth"

AMINE_CV = """\
AMINE KADDOUR
Data Science Master's student | Algiers, Algeria
amine.kaddour@example.org | +213 555 000 111

SUMMARY
Data science master's student specialising in machine learning, seeking
remote ML internships and research opportunities.

SKILLS
Python, SQL, Pandas, NumPy, PyTorch, scikit-learn, Machine Learning,
Deep Learning, Git, Linux

LANGUAGES
Arabic (native), French (fluent), English (B2)

EDUCATION
MSc Data Science, USTHB Algiers, 2025 - 2027
BSc Computer Science, USTHB Algiers, 2022 - 2025

EXPERIENCE
Machine Learning Intern, NAMOCO Analytics, Jun - Sep 2025
Built a churn prediction model with Python and scikit-learn.
Data Analyst (Freelance), 2024 - present
Dashboards and reporting for SME clients.

PROJECTS
Churn-Prediction-Toolkit - https://github.com/aminek/churn-toolkit
Python, scikit-learn, pandas.

CERTIFICATES
Deep Learning Bootcamp, ALX Africa, March 2025, 120 hours

INTERESTS
NLP, computer vision, AI research
"""

SARA_CV = """\
SARA BELKACEM
Frontend Developer | Oran, Algeria
sara.belkacem@example.org

SKILLS
React, TypeScript, Next.js, CSS, Figma, React Native

LANGUAGES
French (native), English (B2)

EDUCATION
BSc Web Development, ENSMAD Oran, 2021 - 2024

EXPERIENCE
Frontend Developer, Startup Studio Oran, 2024 - present
React and React Native apps for logistics clients.

PROJECTS
Logistics Tracker - https://github.com/sarab/logistics-tracker
Design System Starter - https://github.com/sarab/ds-starter
"""

CERTIFICATE = """\
CERTIFICATE OF COMPLETION

This certifies that

AMINE KADDOUR

has successfully completed the

Deep Learning Bootcamp

Issued by ALX Africa - March 2025 - 120 hours
Credential: https://credentials.alx.example/amine-dl-2025
"""


def write_text_pdf(path: Path, text: str, title_size: float = 12.0) -> None:
    """Create a text-based A4 PDF (selectable text, no image)."""
    doc = pymupdf.open()
    page = doc.new_page()
    rect = pymupdf.Rect(50, 50, 545, 792)
    page.insert_textbox(rect, text, fontsize=10, fontname="helv", lineheight=1.35)
    doc.save(path)
    doc.close()


def write_scanned_pdf(path: Path, text: str) -> None:
    """Create an image-only PDF: the text is rendered to pixels (no text layer),
    which forces the OCR path in documents/reader.py."""
    scratch = pymupdf.open()
    page = scratch.new_page()
    page.insert_textbox(
        pymupdf.Rect(50, 50, 545, 792), text, fontsize=10.5, fontname="helv", lineheight=1.4
    )
    pixmap = page.get_pixmap(dpi=150)
    scratch.close()

    doc = pymupdf.open()
    pdf_page = doc.new_page(width=pixmap.width, height=pixmap.height)
    pdf_page.insert_image(pdf_page.rect, pixmap=pixmap)
    doc.save(path)
    doc.close()


def write_portfolio_docx(path: Path) -> None:
    """Create Sara's portfolio as a DOCX with projects and links."""
    doc = Document()
    doc.add_heading("Portfolio - Sara Belkacem", level=1)
    doc.add_paragraph("Frontend developer, Oran, Algeria - sara.belkacem@example.org")

    doc.add_heading("Logistics Tracker", level=2)
    doc.add_paragraph(
        "React Native app tracking delivery fleets for a logistics startup. "
        "Offline-first maps and status sync."
    )
    doc.add_paragraph("Stack: React Native, TypeScript, GraphQL")
    doc.add_paragraph("https://github.com/sarab/logistics-tracker")

    doc.add_heading("Design System Starter", level=2)
    doc.add_paragraph("Open-source accessible component library used by three client projects.")
    doc.add_paragraph("Stack: Next.js, TypeScript, CSS modules")
    doc.add_paragraph("https://github.com/sarab/ds-starter")

    doc.save(path)


def write_ground_truth() -> None:
    """One JSON per sample describing exactly what an extractor should find."""
    truths = {
        "cv_amine.json": {
            "full_name": "Amine Kaddour",
            "headline": "Data Science Master's student",
            "email": "amine.kaddour@example.org",
            "country": "Algeria",
            "skills": [
                "python",
                "sql",
                "pandas",
                "numpy",
                "pytorch",
                "scikit-learn",
                "machine learning",
                "deep learning",
                "git",
                "linux",
            ],
            "languages": ["arabic", "french", "english"],
            "education": [
                {
                    "institution": "USTHB",
                    "degree": "MSc",
                    "field": "Data Science",
                    "start_date": "2025",
                    "end_date": "2027",
                },
                {
                    "institution": "USTHB",
                    "degree": "BSc",
                    "field": "Computer Science",
                    "start_date": "2022",
                    "end_date": "2025",
                },
            ],
            "experience": [
                {
                    "organization": "NAMOCO Analytics",
                    "title": "Machine Learning Intern",
                    "start_date": "2025-06",
                    "end_date": "2025-09",
                },
                {
                    "organization": "Freelance",
                    "title": "Data Analyst",
                    "start_date": "2024",
                    "end_date": None,
                },
            ],
            "certificates": [{"title": "Deep Learning Bootcamp", "issuer": "ALX Africa"}],
            "projects": [
                {
                    "name": "Churn-Prediction-Toolkit",
                    "url": "https://github.com/aminek/churn-toolkit",
                }
            ],
            "interests": ["nlp", "computer vision", "ai research"],
        },
        "cv_sara_scanned.json": {
            "full_name": "Sara Belkacem",
            "headline": "Frontend Developer",
            "email": "sara.belkacem@example.org",
            "country": "Algeria",
            "skills": ["react", "typescript", "next.js", "css", "figma", "react native"],
            "languages": ["french", "english"],
            "education": [
                {
                    "institution": "ENSMAD",
                    "degree": "BSc",
                    "field": "Web Development",
                    "start_date": "2021",
                    "end_date": "2024",
                }
            ],
            "experience": [
                {
                    "organization": "Startup Studio Oran",
                    "title": "Frontend Developer",
                    "start_date": "2024",
                }
            ],
            "projects": [
                {"name": "Logistics Tracker", "url": "https://github.com/sarab/logistics-tracker"},
                {"name": "Design System Starter", "url": "https://github.com/sarab/ds-starter"},
            ],
        },
        "certificate_alx_amine.json": {
            "title": "Deep Learning Bootcamp",
            "issuer": "ALX Africa",
            "holder": "Amine Kaddour",
            "issued_on": "March 2025",
            "hours": 120,
            "credential_url": "https://credentials.alx.example/amine-dl-2025",
        },
        "portfolio_sara.json": {
            "full_name": "Sara Belkacem",
            "email": "sara.belkacem@example.org",
            "projects": [
                {
                    "name": "Logistics Tracker",
                    "url": "https://github.com/sarab/logistics-tracker",
                    "technologies": ["react native", "typescript", "graphql"],
                },
                {
                    "name": "Design System Starter",
                    "url": "https://github.com/sarab/ds-starter",
                    "technologies": ["next.js", "typescript", "css modules"],
                },
            ],
        },
    }
    for name, truth in truths.items():
        (GROUND_TRUTH / name).write_text(json.dumps(truth, indent=2) + "\n", encoding="utf-8")
    print(f"wrote {len(truths)} ground-truth files")


def main() -> None:
    SAMPLES.mkdir(parents=True, exist_ok=True)
    GROUND_TRUTH.mkdir(parents=True, exist_ok=True)

    write_text_pdf(SAMPLES / "cv_amine.pdf", AMINE_CV)
    write_scanned_pdf(SAMPLES / "cv_sara_scanned.pdf", SARA_CV)
    write_text_pdf(SAMPLES / "certificate_alx_amine.pdf", CERTIFICATE)

    doc = Document()
    for line in AMINE_CV.splitlines():
        doc.add_paragraph(line)
    doc.save(SAMPLES / "cv_amine.docx")

    write_portfolio_docx(SAMPLES / "portfolio_sara.docx")
    write_ground_truth()
    print(
        "wrote cv_amine.pdf, cv_amine.docx, cv_sara_scanned.pdf, "
        "certificate_alx_amine.pdf, portfolio_sara.docx"
    )


if __name__ == "__main__":
    main()
